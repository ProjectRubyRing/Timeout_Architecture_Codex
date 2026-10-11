"""Portable visual review using actual Meiryo UI metrics; not a native Excel renderer."""
import json
import math
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
BOOK=ROOT/'output'/'AWS_Timeout_Design_2026-10-11.xlsx'
OUT=ROOT/'validation'

@lru_cache(None)
def font(points,bold=False):
    path='C:/Windows/Fonts/meiryob.ttc' if bold else 'C:/Windows/Fonts/meiryo.ttc'
    return ImageFont.truetype(path,round(points*96/72),index=2)

def color(clr,default):
    if clr is not None and clr.type=='rgb' and isinstance(clr.rgb,str):
        return '#'+clr.rgb[-6:]
    return default

def display(cell):
    val=cell.value
    if val is None:return ''
    if isinstance(val,(int,float)):
        if cell.number_format=='0%':return str(round(val*100))+'%'
        if cell.number_format=='0.0':return f'{val:.1f}'
    return str(val)

def wrap(text,f,width):
    lines=[]
    for paragraph in text.split('\n'):
        line=''
        for ch in paragraph:
            if line and f.getlength(line+ch)>width:
                lines.append(line);line=''
            line+=ch
        lines.append(line)
    return lines

def layout(ws):
    xs=[0]
    for c in range(1,ws.max_column+1):
        width=8.43
        for dim in ws.column_dimensions.values():
            if (dim.min or 0)<=c<=(dim.max or 0):
                width=dim.width
                break
        xs.append(xs[-1]+round(width*7+5))
    ys=[0]
    for r in range(1,ws.max_row+1):
        ys.append(ys[-1]+round((ws.row_dimensions[r].height or 24)*96/72))
    merges={}
    skip=set()
    for area in ws.merged_cells.ranges:
        merges[(area.min_row,area.min_col)]=(area.max_row,area.max_col)
        for r in range(area.min_row,area.max_row+1):
            for c in range(area.min_col,area.max_col+1):
                if (r,c)!=(area.min_row,area.min_col):skip.add((r,c))
    return xs,ys,merges,skip

def main():
    OUT.mkdir(exist_ok=True)
    book=load_workbook(BOOK,data_only=True)
    overflow=[]
    for ws in book:
        xs,ys,merges,skip=layout(ws)
        for row in ws:
            for cell in row:
                r,c=cell.row,cell.column
                if cell.value is None or (r,c) in skip:continue
                r2,c2=merges.get((r,c),(r,c))
                width=xs[c2]-xs[c-1]-12
                height=ys[r2]-ys[r-1]-10
                f=font(cell.font.sz or 10.5,bool(cell.font.bold))
                lines=wrap(display(cell),f,width)
                lh=round((cell.font.sz or 10.5)*96/72*1.28)
                need=len(lines)*lh
                if need>height+3:
                    overflow.append({'sheet':ws.title,'cell':cell.coordinate,'need_px':need,'available_px':height,'text':display(cell)[:80]})
    # Preview representative ranges. They are reference renders, not screenshots of Excel.
    cases=[('01_結論と前提',1,10),('02a_構成図',1,29),('04_設定値一覧',1,11),
           ('06_JBossとJavaの動作',1,10),('13_予算計算',1,22),('15_設定と確認例',1,16)]
    paths=[]
    for name,start,end in cases:
        ws=book[name]
        end=min(end,ws.max_row)
        xs,ys,merges,skip=layout(ws)
        y0=ys[start-1]
        img=Image.new('RGB',(xs[-1]+40,ys[end]-y0+65),'#CCCCCC')
        draw=ImageDraw.Draw(img)
        draw.text((20,5),'Meiryo UI / XLSXセルからの参考描画（Excel実画面ではありません）',font=font(9),fill='#222222')
        for row in ws.iter_rows(min_row=start,max_row=end):
            for cell in row:
                r,c=cell.row,cell.column
                if (r,c) in skip:continue
                r2,c2=merges.get((r,c),(r,c))
                if r2>end:continue
                x,y=xs[c-1]+20,ys[r-1]-y0+35
                right,bottom=xs[c2]+20,ys[r2]-y0+35
                bg=color(cell.fill.fgColor,'#FFFFFF') if cell.fill.patternType else '#FFFFFF'
                draw.rectangle((x,y,right,bottom),fill=bg)
                if cell.border.bottom.style or cell.border.left.style:
                    draw.rectangle((x,y,right,bottom),outline='#CFCFCF',width=1)
                text=display(cell)
                if not text:continue
                f=font(cell.font.sz or 10.5,bool(cell.font.bold))
                lines=wrap(text,f,right-x-12)
                lh=round((cell.font.sz or 10.5)*96/72*1.28)
                ty=y+5
                if cell.alignment.vertical=='center':ty=y+(bottom-y-len(lines)*lh)/2
                for line in lines:
                    tx=x+6
                    if cell.alignment.horizontal=='center':tx=x+(right-x-f.getlength(line))/2
                    elif cell.alignment.horizontal=='right':tx=right-6-f.getlength(line)
                    draw.text((tx,ty),line,font=f,fill=color(cell.font.color,'#222222'))
                    ty+=lh
        path=OUT/(name+'_preview.png')
        img.save(path)
        paths.append(str(path))
    result={'render_type':'Pillow using XLSX cells and installed Meiryo UI fonts; not native Excel',
            'overflow_candidates':overflow,'previews':paths}
    (OUT/'layout_checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=True,indent=2))

if __name__=='__main__':main()
