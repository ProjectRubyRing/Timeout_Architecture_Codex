"""Artifact checks; does not claim to validate an AWS deployment."""
import json
import re
import zipfile
from pathlib import Path
from hashlib import sha256

from openpyxl import load_workbook
from report_data import PARAMS, SOURCES, SECTIONS, EXAMPLES

ROOT = Path(__file__).resolve().parents[1]
STEM = 'AWS_Timeout_Design_2026-10-11'
XLSX = ROOT / 'output' / (STEM + '.xlsx')
MD = ROOT / 'output' / (STEM + '.md')
VALIDATION = ROOT / 'validation'
VALIDATION.mkdir(exist_ok=True)

def main():
    errors=[]
    book=load_workbook(XLSX,data_only=False)
    cached=load_workbook(XLSX,data_only=True)
    md=MD.read_text(encoding='utf-8')
    sources={s[0]:s for s in SOURCES}
    fonts=set()
    populated=0
    formulas={}
    for ws in book:
        if ws.sheet_view.showGridLines is not False:
            errors.append(f'{ws.title}: visible grid')
        if not ws.freeze_panes:
            errors.append(f'{ws.title}: freeze missing')
        if not ws.print_area:
            errors.append(f'{ws.title}: print area missing')
        for row in ws:
            for cell in row:
                if cell.value is None:continue
                populated+=1
                fonts.add(cell.font.name)
                if cell.font.name != 'Meiryo UI':
                    errors.append(f'{ws.title}/{cell.coordinate}: font {cell.font.name}')
                if cell.alignment.wrap_text is not True:
                    errors.append(f'{ws.title}/{cell.coordinate}: no wrap')
                if cell.data_type=='f':
                    cv=cached[ws.title][cell.coordinate].value
                    formulas[f'{ws.title}!{cell.coordinate}']={'formula':cell.value,'cached':cv}
                    if cv is None or str(cv).startswith(('#REF!','#VALUE!','#DIV/0!','#NAME?')):
                        errors.append(f'{ws.title}/{cell.coordinate}: invalid cache')
                if '\ufffd' in str(cell.value):
                    errors.append(f'{ws.title}/{cell.coordinate}: Unicode replacement char')
                for color in [cell.font.color,cell.fill.fgColor]:
                    if color and color.type=='rgb' and isinstance(color.rgb,str):
                        rgb=color.rgb[-6:]
                        if rgb[0:2]!=rgb[2:4] or rgb[2:4]!=rgb[4:6]:
                            errors.append(f'{ws.title}/{cell.coordinate}: non-monochrome {rgb}')
                if cell.hyperlink and cell.hyperlink.location:
                    location=cell.hyperlink.location
                    match=re.match(r"'([^']+)'!(.+)",location)
                    if not match or match[1] not in book.sheetnames:
                        errors.append(f'Broken internal link: {location}')
        for rd in ws.row_dimensions.values():
            if rd.height and rd.height>409:
                errors.append(f'{ws.title}: excessive height {rd.height}')
    expected_caches={'B13':25,'B19':4,'B20':6,'B22':'予算内','B27':2.5,'B28':450,
                     'B31':'余白あり','B37':700,'B40':585,'B41':'枠内'}
    for cell,val in expected_caches.items():
        actual=cached['13_予算計算'][cell].value
        if actual!=val:errors.append(f'Cached calculation {cell}: {actual} != {val}')
    index=book['04_設定値一覧']
    for r,p in enumerate(PARAMS,6):
        wanted=[p[k] for k in ['id','key','kind','default','online','batch','scope','refs']]
        actual=[index.cell(r,c).value for c in range(1,9)]
        actual[-1]=' '.join((actual[-1] or '').split())
        if actual!=wanted:errors.append(f'Parameter mismatch: {p["id"]}')
        for k in ['key','default','online','batch','start','effect','reason','caveat']:
            if p[k] not in md:errors.append(f'Markdown missing {p["id"]}.{k}')
        for sid in p['refs'].split():
            if sid not in sources:errors.append(f'Unknown source {sid}')
    for sec in SECTIONS:
        for label,text,refs in sec['items']:
            if text not in md:errors.append(f'Missing narrative {label}')
    for e in EXAMPLES:
        if e['code'] not in md:errors.append(f'Missing example {e["title"]}')
    for sid,_,_,url,_ in SOURCES:
        if url not in md:errors.append(f'Missing source URL {sid}')
    if len(book.sheetnames)!=23:errors.append('Unexpected sheet count')
    if '\ufffd' in md:errors.append('Unicode replacement char in MD')
    for path in [*ROOT.joinpath('source').glob('*.py'),MD]:
        for no,line in enumerate(path.read_text('utf-8').splitlines(),1):
            if line.rstrip()!=line:errors.append(f'Trailing whitespace {path.name}:{no}')
    with zipfile.ZipFile(XLSX) as archive:
        if archive.testzip() is not None:errors.append('Bad XLSX ZIP CRC')
        if any('vbaProject' in n for n in archive.namelist()):errors.append('Unexpected macro')
    result={'artifact_checks':'PASS' if not errors else 'FAIL','sheets':len(book.sheetnames),
            'parameters':len(PARAMS),'sources':len(SOURCES),'populated_cells':populated,
            'cell_fonts':sorted(fonts),'formula_caches':formulas,'markdown_characters':len(md),
            'markdown_lines':len(md.splitlines()),'errors':errors,
            'sha256':{p.name:sha256(p.read_bytes()).hexdigest() for p in [XLSX,MD]},
            'runtime_validation':'AWS / JBoss / MySQL integration and load tests NOT RUN',
            'git':'Workspace is not a Git repository; text whitespace and file inventory checked separately.'}
    (VALIDATION/'artifact_checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=True,indent=2))
    if errors:raise SystemExit(1)

if __name__=='__main__':main()
